"""Data storage, querying, and reporting layer for Required Product Requisition Slips.

Provides MongoDB persistence for requisitions and Excel (.xlsx) export generation.
"""
from __future__ import annotations

import io
import logging
from datetime import datetime
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pymongo.database import Database

from config.settings import STORE_CODE_TO_NAME
from db.models import REQUISITIONS, next_id
from src.product_hierarchy import register_custom_category

logger = logging.getLogger(__name__)

STORE_FULL_NAMES = {
    "NM": "LOURDES TEXTILES PVT. LTD. UNIT - CITIMART NEW MARKET",
    "HB": "LOURDES TEXTILES PVT. LTD. UNIT - CITIMART HATIBAGAN",
    "CHW": "LOURDES TEXTILES PVT. LTD. UNIT - CITIMART CHOWRINGHEE",
}


def _format_product_required(sl: int, division: str, section: str, department: str) -> str:
    """Format product required hierarchy string.
    Example: '(1) -> Accessories -> Gift & Novelties -> Baby Accessory'
    """
    div = division.strip() if division else ""
    sec = section.strip() if section else ""
    dep = department.strip() if department else ""
    parts = [p for p in [div, sec, dep] if p]
    hierarchy_str = " -> ".join(parts) if parts else "N/A"
    return f"({sl}) -> {hierarchy_str}"


def create_requisition(
    db: Database,
    store_code: str,
    date_str: str,
    items: List[Dict[str, Any]],
    created_by_user: str,
    created_by_name: str,
    created_by_role: str,
    priority: str = "Normal",
    remarks_general: str = "",
) -> Dict[str, Any]:
    """Create a new product requisition slip with one or more items."""
    req_id = next_id(db, REQUISITIONS)
    now_iso = datetime.now().isoformat()
    store_code = store_code.upper()
    store_name_full = STORE_FULL_NAMES.get(
        store_code, f"LOURDES TEXTILES PVT. LTD. UNIT - CITIMART {STORE_CODE_TO_NAME.get(store_code, store_code)}"
    )

    date_clean = date_str.strip() or datetime.now().strftime("%Y-%m-%d")
    date_compact = date_clean.replace("-", "")
    req_code = f"REQ-{date_compact}-{store_code}-{req_id:04d}"

    formatted_items: List[Dict[str, Any]] = []
    for idx, raw_item in enumerate(items, start=1):
        sl = raw_item.get("sl_no") or idx
        div = str(raw_item.get("division") or "").strip()
        sec = str(raw_item.get("section") or "").strip()
        dep = str(raw_item.get("department") or "").strip()

        prod_required = raw_item.get("product_required") or _format_product_required(sl, div, sec, dep)
        barcode = str(raw_item.get("barcode_details") or raw_item.get("barcode") or "").strip()
        brand = str(raw_item.get("brand") or "").strip()
        
        mrp_raw = raw_item.get("mrp")
        mrp_str = str(mrp_raw).strip() if mrp_raw is not None else ""

        time_req = str(raw_item.get("time_required") or "").strip() or "1 day"
        item_remarks = str(raw_item.get("remarks") or "").strip()

        # Register custom categories to persistent custom category collection
        if div:
            register_custom_category(db, div, sec, dep)

        formatted_items.append({
            "sl_no": sl,
            "division": div,
            "section": sec,
            "department": dep,
            "product_required": prod_required,
            "barcode_details": barcode,
            "brand": brand,
            "mrp": mrp_str,
            "time_required": time_req,
            "remarks": item_remarks,
        })

    doc: Dict[str, Any] = {
        "_id": req_id,
        "req_code": req_code,
        "store_code": store_code,
        "store_name": STORE_CODE_TO_NAME.get(store_code, store_code),
        "store_name_full": store_name_full,
        "date": date_clean,
        "created_at": now_iso,
        "updated_at": now_iso,
        "created_by_user": created_by_user,
        "created_by_name": created_by_name,
        "created_by_role": created_by_role,
        "target_recipient": "Operational Head / Admin",
        "priority": priority,
        "status": "Pending",  # Pending | In Review | Approved | Fulfilled | Rejected
        "remarks_general": remarks_general,
        "items": formatted_items,
        "item_count": len(formatted_items),
        "admin_remarks": "",
        "actioned_by": "",
        "actioned_at": "",
    }

    db[REQUISITIONS].insert_one(doc)
    return doc


def list_requisitions(
    db: Database,
    store_code: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    status: Optional[str] = None,
    search: Optional[str] = None,
    limit: int = 500,
) -> List[Dict[str, Any]]:
    """List requisitions with multi-criteria filtering."""
    query: Dict[str, Any] = {}

    if store_code and store_code.upper() not in ("ALL", ""):
        query["store_code"] = store_code.upper()

    if start_date and end_date:
        query["date"] = {"$gte": start_date, "$lte": end_date}
    elif start_date:
        query["date"] = {"$gte": start_date}
    elif end_date:
        query["date"] = {"$lte": end_date}

    if status and status.upper() not in ("ALL", ""):
        query["status"] = status

    if search and search.strip():
        term = search.strip()
        query["$or"] = [
            {"req_code": {"$regex": term, "$options": "i"}},
            {"store_name_full": {"$regex": term, "$options": "i"}},
            {"created_by_name": {"$regex": term, "$options": "i"}},
            {"remarks_general": {"$regex": term, "$options": "i"}},
            {"items.product_required": {"$regex": term, "$options": "i"}},
            {"items.barcode_details": {"$regex": term, "$options": "i"}},
            {"items.brand": {"$regex": term, "$options": "i"}},
            {"items.department": {"$regex": term, "$options": "i"}},
            {"items.section": {"$regex": term, "$options": "i"}},
            {"items.division": {"$regex": term, "$options": "i"}},
            {"items.remarks": {"$regex": term, "$options": "i"}},
        ]

    cursor = db[REQUISITIONS].find(query).sort("created_at", -1).limit(limit)
    return list(cursor)


def get_requisition_by_id(db: Database, req_id: int) -> Optional[Dict[str, Any]]:
    """Retrieve a single requisition slip by integer ID."""
    return db[REQUISITIONS].find_one({"_id": req_id})


def update_requisition_status(
    db: Database,
    req_id: int,
    status: str,
    admin_remarks: str = "",
    actioned_by: str = "",
) -> Optional[Dict[str, Any]]:
    """Update requisition status and admin comments."""
    now_iso = datetime.now().isoformat()
    update_data: Dict[str, Any] = {
        "status": status,
        "updated_at": now_iso,
        "actioned_at": now_iso,
    }
    if admin_remarks:
        update_data["admin_remarks"] = admin_remarks
    if actioned_by:
        update_data["actioned_by"] = actioned_by

    res = db[REQUISITIONS].find_one_and_update(
        {"_id": req_id},
        {"$set": update_data},
        return_document=openpyxl.load_workbook if False else True,  # placeholder
    )
    # Re-fetch updated doc
    return db[REQUISITIONS].find_one({"_id": req_id})


def delete_requisition(db: Database, req_id: int) -> bool:
    """Delete a requisition slip."""
    res = db[REQUISITIONS].delete_one({"_id": req_id})
    return res.deleted_count > 0


def generate_requisitions_excel(requisitions: List[Dict[str, Any]]) -> bytes:
    """Generate a clean, styled Excel workbook for Requisition Slips.
    
    Columns:
    sl.no | Date | Store Name | Product Required | BARCODE DETAILS DESCRIPTIONS | BRAND | MRP | Time Required | Remarks | Status | Admin Remarks
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Product Requisitions"
    ws.views.sheetView[0].showGridLines = True

    # Palette
    NAVY_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    HEADER_FONT = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    TITLE_FONT = Font(name="Segoe UI", size=15, bold=True, color="1E3A8A")
    SUBTITLE_FONT = Font(name="Segoe UI", size=10, italic=True, color="475569")
    REGULAR_FONT = Font(name="Segoe UI", size=10, color="0F172A")
    BOLD_FONT = Font(name="Segoe UI", size=10, bold=True, color="0F172A")

    THIN_BORDER = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    # Title Block
    ws.merge_cells("A1:K1")
    ws["A1"] = "LOURDES TEXTILES PVT. LTD. — CITIMART DAILY OPERATIONS"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells("A2:K2")
    ws["A2"] = f"Required Product Requisition Report — Generated on {datetime.now().strftime('%d-%b-%Y %I:%M %p')}"
    ws["A2"].font = SUBTITLE_FONT
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")

    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 10
    ws.row_dimensions[4].height = 26

    # Headers
    headers = [
        "Sl.No",
        "Date",
        "Store Name",
        "Product Required",
        "BARCODE DETAILS DESCRIPTIONS",
        "BRAND",
        "MRP (₹)",
        "Time Required",
        "Remarks",
        "Status",
        "Admin Remarks",
    ]

    for col_idx, header_text in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=header_text)
        cell.fill = NAVY_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    current_row = 5
    global_item_counter = 1

    for req in requisitions:
        req_date = req.get("date", "")
        store_full = req.get("store_name_full", req.get("store_name", ""))
        status = req.get("status", "Pending")
        admin_rem = req.get("admin_remarks", "")
        items = req.get("items", [])

        if not items:
            # Requisition without sub-items
            row_data = [
                global_item_counter,
                req_date,
                store_full,
                "-",
                "-",
                "-",
                "-",
                "-",
                req.get("remarks_general", ""),
                status,
                admin_rem,
            ]
            fill = ZEBRA_FILL if global_item_counter % 2 == 0 else PatternFill(fill_type=None)
            for col_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.font = REGULAR_FONT
                cell.border = THIN_BORDER
                if fill.fill_type:
                    cell.fill = fill
                cell.alignment = Alignment(
                    horizontal="center" if col_idx in (1, 2, 7, 8, 10) else "left",
                    vertical="center",
                    wrap_text=True,
                )
            ws.row_dimensions[current_row].height = 22
            current_row += 1
            global_item_counter += 1
        else:
            for item in items:
                row_data = [
                    global_item_counter,
                    req_date,
                    store_full,
                    item.get("product_required", ""),
                    item.get("barcode_details", ""),
                    item.get("brand", ""),
                    item.get("mrp", ""),
                    item.get("time_required", ""),
                    item.get("remarks", "") or req.get("remarks_general", ""),
                    status,
                    admin_rem,
                ]
                fill = ZEBRA_FILL if global_item_counter % 2 == 0 else PatternFill(fill_type=None)
                for col_idx, val in enumerate(row_data, start=1):
                    cell = ws.cell(row=current_row, column=col_idx, value=val)
                    cell.font = BOLD_FONT if col_idx == 4 else REGULAR_FONT
                    cell.border = THIN_BORDER
                    if fill.fill_type:
                        cell.fill = fill
                    cell.alignment = Alignment(
                        horizontal="center" if col_idx in (1, 2, 7, 8, 10) else "left",
                        vertical="center",
                        wrap_text=True,
                    )
                ws.row_dimensions[current_row].height = 24
                current_row += 1
                global_item_counter += 1

    # Auto-fit column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            if cell.row < 4:
                continue
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # Set specific generous widths for wide columns
    ws.column_dimensions["A"].width = 8   # Sl.No
    ws.column_dimensions["B"].width = 13  # Date
    ws.column_dimensions["C"].width = 35  # Store Name
    ws.column_dimensions["D"].width = 45  # Product Required
    ws.column_dimensions["E"].width = 32  # Barcode
    ws.column_dimensions["F"].width = 18  # Brand
    ws.column_dimensions["G"].width = 14  # MRP
    ws.column_dimensions["H"].width = 16  # Time Required
    ws.column_dimensions["I"].width = 28  # Remarks
    ws.column_dimensions["J"].width = 14  # Status
    ws.column_dimensions["K"].width = 28  # Admin Remarks

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
